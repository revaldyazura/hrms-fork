from __future__ import annotations

from datetime import datetime
from io import BytesIO
import tempfile
from typing import Any, Dict, Optional

import frappe

from hrms.integrations.telegram_bot import utils as telegram_utils
from hrms.integrations.telegram_bot.aduan import aduan

DEFAULT_BULK_COMMAND = "aduan_bulk"
TEMPLATE_BASENAME = "aduan_bulk_template.xlsx"


def _bulk_text(key: str, default: str, context: Optional[dict] = None) -> str:
    """Render a templated text via telegram_aduan_response_texts (if configured)."""

    return (telegram_utils._render_response_text(key, default, context or {}) or "").strip()


def is_bulk_command(command_token: str) -> bool:
    name = telegram_utils.command_for_menu(command_token)
    if not name:
        return False
    if name.lower().startswith(DEFAULT_BULK_COMMAND):
        print(f"Command {command_token} is recognized as bulk command")
        return True
    return False


def template_file_path() -> str:
    return frappe.get_app_path(
        "hrms",
        "integrations",
        "telegram_bot",
        "aduan",
        "templates",
        TEMPLATE_BASENAME,
    )


def handle_aduan_bulk_with_report(
    bot: Any,
    message: Any,
    command_token: str,
    *,
    list_limit: int = 10,
) -> tuple[str, Optional[str]]:
    """Like `handle_aduan_bulk`, but optionally generates a .txt report.

    Returns:
    - message_text: str (summary suitable for chat)
    - report_file_path: Optional[str] (path to a generated .txt with full details)
      Generated only when successes or failures exceed `list_limit`.
    """

    cmd_norm = telegram_utils.normalize_command_token(command_token, default="/aduan_bulk")
    base_context = {
        "cmd": telegram_utils._escape_html(cmd_norm)
    }

    doc = getattr(message, "document", None)
    if not doc:
        raise frappe.ValidationError(
            _bulk_text(
                "aduan_bulk_err_no_document",
                "Mohon lampirkan file Excel (.xlsx)",
                base_context,
            )
        )

    filename = (getattr(doc, "file_name", None) or "").strip()
    if not filename.lower().endswith(".xlsx"):
        raise frappe.ValidationError(
            _bulk_text(
                "aduan_bulk_err_invalid_extension",
                "File harus berformat .xlsx",
                {**base_context, "filename": telegram_utils._escape_html(filename)},
            )
        )

    file_id = getattr(doc, "file_id", None)
    if not file_id:
        raise frappe.ValidationError(
            _bulk_text(
                "aduan_bulk_err_missing_file_id",
                "File tidak ditemukan",
                base_context,
            )
        )

    try:
        print("Downloading attached file from Telegram...")
        tg_file = bot.get_file(message.document.file_id)
        file_bytes = bot.download_file(tg_file.file_path)
    except Exception as e:
        raise frappe.ValidationError(
            _bulk_text(
                "aduan_bulk_err_download_failed",
                "Gagal mengunduh file: {error}",
                {**base_context, "error": telegram_utils._escape_html(str(e))},
            )
        )

    rows = _read_excel_rows(file_bytes)
    if not rows:
        raise frappe.ValidationError(
            _bulk_text(
                "aduan_bulk_err_excel_empty",
                "File Excel kosong atau tidak memiliki data",
                base_context,
            )
        )
    chat_id = telegram_utils._coerce_int(getattr(getattr(message, "chat", None), "id", None))
    thread_id = telegram_utils._coerce_int(getattr(message, "message_thread_id", None))
 
    mapping_row = telegram_utils._maintask_mapping_row_for_message(chat_id, thread_id)

    settings = telegram_utils._conf_default_subtask_settings(message)
    base = (frappe.conf.get("telegram_aduan_site") or "hris.ebdesk.com/app/subtask/").strip()

    ok: list[dict[str, str]] = []
    failed: list[dict[str, str]] = []

    for row in rows:
        row_no = int(row.get("_row") or 0)
        fields = {k: v for k, v in row.items() if not k.startswith("_")}
        freeform = (row.get("_freeform") or "").strip()

        missing = aduan._missing_required_fields(fields)
        if missing:
            failed.append(
                {
                    "row": str(row_no),
                    "error": "Missing: " + ", ".join(missing),
                }
            )
            continue

        try:
            pics: list[str] = []
            issue_label = ""

            if mapping_row:
                issue_docname, issue_label, pics = telegram_utils._resolve_issue_type_from_maintask_mapping(
                    fields.get("issue_type") or "",
                    mapping_row,
                    settings.get("maintask") or "",
                    settings.get("maintask_name") or "",
                )
            else:
                issue_docname, issue_label = telegram_utils._resolve_issue_type(
                    fields.get("issue_type") or "",
                    settings.get("maintask") or "",
                    settings.get("maintask_name") or "",
                )
                pics = telegram_utils._pics_for_issue_user_input(fields.get("issue_type") or "")

            subtask_id = _create_subtask_from_bulk(fields, message, settings, issue_docname, freeform=freeform)

            subtask_url = (base + subtask_id) if base else subtask_id
            subtask_link = telegram_utils._format_hyperlink(subtask_id, subtask_url) if base else subtask_id
            pic_text = ", ".join(pics) if pics else ""

            ok.append(
                {
                    "row": str(row_no),
                    "subtask_id": subtask_id,
                    "subtask_url": subtask_url,
                    "subtask_link": subtask_link,
                    "issue": telegram_utils._escape_html(issue_label or ""),
                    "pic": telegram_utils._escape_html(pic_text or ""),
                }
            )
        except Exception as e:
            failed.append({"row": str(row_no), "error": str(e)})
            try:
                frappe.db.rollback()
                try:
                    frappe.db.value_cache.clear()
                except Exception:
                    pass
            except Exception:
                pass

    summary_msg = _build_summary_message(
        ok,
        failed,
        list_limit=list_limit,
        context=base_context,
    )

    report_path: Optional[str] = None
    if len(ok) > list_limit or len(failed) > list_limit:
        report_text = _build_full_report_text(ok, failed, context=base_context)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        # Use a temp file; telegram_aduan_bot will send it via _send(... send_document=True ...)
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            suffix=".txt",
            prefix=f"aduan_bulk_report_{ts}_",
            delete=False,
        ) as f:
            f.write(report_text)
            report_path = f.name

    return summary_msg, report_path


def _build_summary_message(
    ok: list[dict[str, str]],
    failed: list[dict[str, str]],
    *,
    list_limit: int,
    context: Optional[dict] = None,
) -> str:
    ok_shown = ok[:list_limit]
    failed_shown = failed[:list_limit]

    # Keep summary_body as "dynamic data only" so all static headings can live in config templates.
    # Tab-separated to avoid embedding static labels like "Row" in code.
    ok_items = "\n".join(
        [
            "\t".join(
                [
                    (item.get("row") or "").strip(),
                    (item.get("subtask_link") or "").strip(),
                    (item.get("issue") or "").strip(),
                    (item.get("pic") or "").strip(),
                ]
            ).rstrip()
            for item in ok_shown
        ]
    ).strip()

    failed_items = "\n".join(
        [
            "\t".join(
                [
                    (item.get("row") or "").strip(),
                    telegram_utils._escape_html((item.get("error") or "").strip()),
                ]
            ).rstrip()
            for item in failed_shown
        ]
    ).strip()

    summary_body = "\n".join([x for x in (ok_items, failed_items) if x]).strip()

    has_report = len(ok) > list_limit or len(failed) > list_limit
    ctx = {
        **(context or {}),
        "ok_count": str(len(ok)),
        "fail_count": str(len(failed)),
        "ok_shown_count": str(len(ok_shown)),
        "fail_shown_count": str(len(failed_shown)),
        "ok_remaining_count": str(max(0, len(ok) - len(ok_shown))),
        "fail_remaining_count": str(max(0, len(failed) - len(failed_shown))),
        "list_limit": str(list_limit),
        "has_report": "ℹ️ Detail lengkap dikirim sebagai file .txt" if has_report else "ℹ️ Menampilkan semua hasil",
        "summary_body": summary_body,
        "ok_items": ok_items,
        "failed_items": failed_items,
    }

    default = f"✅ Aduan bulk selesai: {len(ok)} sukses, {len(failed)} gagal\n\n" + summary_body
    msg = _bulk_text("aduan_bulk_summary", default.strip(), ctx)
    if len(msg) > 3800:
        msg = msg[:3800].rstrip() + "\n..."
    return msg


def _build_full_report_text(
    ok: list[dict[str, str]],
    failed: list[dict[str, str]],
    *,
    context: Optional[dict] = None,
) -> str:
    # Keep report_body as "dynamic data only" so static headings live in config templates.
    ok_items_full = "\n".join(
        [
            "\t".join(
                [
                    (item.get("row") or "").strip(),
                    (item.get("subtask_id") or "").strip(),
                    (item.get("subtask_url") or "").strip(),
                    (item.get("issue") or "").strip(),
                    (item.get("pic") or "").strip(),
                ]
            ).rstrip()
            for item in (ok or [])
        ]
    ).strip()

    failed_items_full = "\n".join(
        [
            "\t".join(
                [
                    (item.get("row") or "").strip(),
                    (item.get("error") or "").strip(),
                ]
            ).rstrip()
            for item in (failed or [])
        ]
    ).strip()

    report_body = "\n".join([x for x in (ok_items_full, failed_items_full) if x]).rstrip() + "\n"
    ctx = {
        **(context or {}),
        "ok_count": str(len(ok)),
        "fail_count": str(len(failed)),
        "report_body": report_body,
        "ok_items": ok_items_full,
        "failed_items": failed_items_full,
    }
    default = report_body
    return _bulk_text("aduan_bulk_report", default, ctx)


def handle_aduan_bulk(bot: Any, message: Any, command_token: str) -> str:
    """Process an /aduan_bulk message containing an .xlsx document."""

    # Backwards-compatible wrapper: return only the summary string.
    msg, _report_path = handle_aduan_bulk_with_report(bot, message, command_token=command_token)
    return msg


def _read_excel_rows(file_bytes: bytes) -> list[dict[str, str]]:
    """Read .xlsx bytes and return rows as field dicts.

    Expected headers similar to /aduan fields.
    """

    try:
        print("Importing openpyxl for reading Excel file...")
        from openpyxl import load_workbook
    except Exception as e:
        raise frappe.ValidationError(
            _bulk_text(
                "aduan_bulk_err_openpyxl_missing",
                "openpyxl belum tersedia: {error}",
                {"error": str(e)},
            )
        )

    wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return []

    header_map: dict[int, str] = {}
    extra_header_by_idx: dict[int, str] = {}
    details_idx: Optional[int] = None
    for idx, raw in enumerate(header_row or []):
        key = _normalize_header(raw)
        if not key:
            continue
        normalized = _HEADER_ALIASES.get(key)
        if normalized:
            header_map[idx] = normalized
            if normalized == "details" and details_idx is None:
                details_idx = idx

        # Any column to the right of "Details" should be treated as extra description.
        # This allows adding new columns without touching _HEADER_ALIASES.
        if details_idx is not None and idx > details_idx:
            hdr = "" if raw in (None, "") else str(raw).strip()
            if hdr:
                extra_header_by_idx[idx] = hdr

    if not header_map:
        raise frappe.ValidationError(
            _bulk_text(
                "aduan_bulk_err_header_unknown",
                "Header Excel tidak dikenali",
            )
        )

    out: list[dict[str, str]] = []
    excel_row_no = 1
    for excel_row_no, row_values in enumerate(rows_iter, start=2):
        if not row_values:
            continue

        row: dict[str, str] = {"_row": str(excel_row_no)}
        any_value = False
        freeform_lines: list[str] = []

        for idx, raw_val in enumerate(row_values):
            val = "" if raw_val in (None, "") else str(raw_val).strip()
            val = telegram_utils._clean_field_value(val)
            if val:
                any_value = True

            # 1) Known columns (mapped via _HEADER_ALIASES)
            field = header_map.get(idx)
            if field:
                if field == "issue_type":
                    row["issue_type"] = val
                elif field == "link_dashboard":
                    row["link_dashboard"] = val
                elif field == "link_maps":
                    row["link_maps"] = val
                else:
                    row[field] = val
                continue

            # 2) Extra columns to the right of Details => freeform description
            hdr = extra_header_by_idx.get(idx)
            if hdr and val:
                freeform_lines.append(f"{hdr}: {val}")

        if freeform_lines:
            row["_freeform"] = "\n".join(freeform_lines).strip()

        if any_value:
            out.append(row)

    return out


def _normalize_header(value: object) -> str:
    if value in (None, ""):
        return ""
    s = str(value).strip().lower()
    s = " ".join(s.split())
    return s


_HEADER_ALIASES: dict[str, str] = {
    "type": "type",
    "priority": "priority",
    "issues": "issue_type",
    "issue type": "issue_type",
    "subject": "subject",
    "details": "details",
    "link dashboard": "link_dashboard",
    "link maps": "link_maps",
    "workspace": "workspace",
}


def _create_subtask_from_bulk(
    fields: Dict[str, str],
    message: Any,
    settings: Dict[str, Any],
    issue_docname: str,
    *,
    freeform: str = "",
) -> str:
    maintask = settings["maintask"]
    tasks = settings["tasks"]
    owner = settings["owner"]
    pic_subtask = settings["pic_subtask"]

    subtask_name = (fields.get("subject") or "").strip()
    if not subtask_name:
        raise frappe.ValidationError("Subject is required")

    priority = telegram_utils._normalize_priority(fields.get("priority"))
    description = telegram_utils._build_description(fields, freeform or "", message)

    doc = frappe.get_doc(
        {
            "doctype": "SubTask",
            "maintask": maintask,
            "tasks": tasks,
            "subtask_name": subtask_name,
            "description": description,
            "requestor": telegram_utils._telegram_user_label(message),
            "created_by": telegram_utils._telegram_user_label(message),
            "priority": priority,
            "status": "Open",
            "pic_subtask": pic_subtask,
            "target_time": int(settings["target_time"]),
            "unit_target_time": settings["unit_target_time"],
            "value": str(settings["value"]),
            "owner": owner,
        }
    )

    if fields.get("type"):
        type_name = telegram_utils._resolve_subtask_type(fields["type"])
        doc.append("type", {"subtask_type": type_name})

    doc.append("issues_type", {"issue": issue_docname})

    aduan._insert_with_owner(doc, owner)
    frappe.db.commit()
    return doc.name
